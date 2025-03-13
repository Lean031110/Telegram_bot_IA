import sqlite3
import nltk
from collections import Counter
from nltk.corpus import stopwords
from nltk.stem import WordNetLemmatizer
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.naive_bayes import MultinomialNB
from sklearn.pipeline import make_pipeline
import joblib
import pandas as pd
from docx import Document
import openpyxl

DB_NAME = "chat_data.db"

nltk.download('punkt')
nltk.download('stopwords')
nltk.download('wordnet')
nltk.download('omw-1.4')

def analyze_chat():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute("SELECT message FROM messages")
    messages = [row[0] for row in cursor.fetchall()]

    conn.close()

    # Tokenización y lematización
    lemmatizer = WordNetLemmatizer()
    stop_words = set(stopwords.words('spanish'))

    words = nltk.word_tokenize(" ".join(messages))
    words = [lemmatizer.lemmatize(word.lower()) for word in words if word.isalnum() and word.lower() not in stop_words]

    word_freq = Counter(words)
    common_words = word_freq.most_common(10)

    return common_words

def train_intent_classifier():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute("SELECT message, intent FROM messages")
    data = cursor.fetchall()

    conn.close()

    df = pd.DataFrame(data, columns=["message", "intent"])

    # Llenar valores nulos (None) con una intención predeterminada
    df["intent"].fillna("desconocido", inplace=True)

    model = make_pipeline(TfidfVectorizer(), MultinomialNB())
    model.fit(df["message"], df["intent"])

    joblib.dump(model, "intent_classifier.pkl")

    return model

def classify_intent(message):
    model = joblib.load("intent_classifier.pkl")
    return model.predict([message])[0]

def train_sentiment_analyzer():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute("SELECT message, sentiment FROM messages")
    data = cursor.fetchall()

    conn.close()

    if not data:
        return None

    df = pd.DataFrame(data, columns=["message", "sentiment"])

    # Llenar valores nulos (None) con un sentimiento predeterminado
    df["sentiment"].fillna("neutral", inplace=True)

    model = make_pipeline(TfidfVectorizer(), MultinomialNB())
    model.fit(df["message"], df["sentiment"])

    joblib.dump(model, "sentiment_analyzer.pkl")

    return model

def analyze_sentiment(message):
    model = joblib.load("sentiment_analyzer.pkl")
    return model.predict([message])[0]

def retrain_models():
    train_intent_classifier()
    train_sentiment_analyzer()

def process_word_document(file_path):
    document = Document(file_path)
    messages = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            parts = text.split('|')
            if len(parts) == 2:
                messages.append((parts[0].strip(), parts[1].strip()))
    return messages

def process_excel_document(file_path):
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    messages = []
    for row in sheet.iter_rows(values_only=True):
        if len(row) == 2:
            message, intent = row
            messages.append((str(message).strip(), str(intent).strip()))
    return messages

def add_messages_to_db(messages):
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    for message, intent in messages:
        cursor.execute("INSERT INTO messages (message, intent) VALUES (?, ?)", (message, intent))

    conn.commit()
    conn.close()

def generate_response(intent):
    responses = {
        "saludo": ["¡Hola! ¿Cómo estás?", "Buenos días", "Buenas tardes", "Buenas noches"],
        "nombre_bot": ["Soy un bot de Telegram con inteligencia artificial.", "Mi nombre es BotIA."],
        "funcionalidades": ["Puedo ayudarte con varias tareas, responder preguntas, y moderar el chat.", "Mis capacidades incluyen responder preguntas, moderar el chat, y proporcionar información útil."],
        "proposito": ["Mi propósito es ayudarte a gestionar este chat y proporcionar información útil.", "Fui creado para asistirte en diversas tareas y proporcionar información."],
        "funcionamiento": ["Analizo los mensajes y respondo según mi programación y entrenamiento.", "Funciono mediante algoritmos de inteligencia artificial que procesan y analizan la información."],
        "chiste": ["¿Por qué los programadores prefieren el teclado? Porque les da más control.", "¿Qué hace una abeja en el gimnasio? ¡Zum-ba!"],
        "hora": ["No tengo acceso a la hora actual.", "Lo siento, pero no puedo decirte la hora."],
        "fecha": ["No tengo acceso a la fecha actual.", "Lo siento, pero no puedo decirte la fecha."],
        "clima": ["No tengo acceso a información del clima.", "Lo siento, pero no puedo proporcionarte el pronóstico del tiempo."],
        "desconocido": ["Lo siento, no tengo una respuesta para eso.", "No entiendo tu pregunta, ¿puedes reformularla?"]
    }
    return responses.get(intent, ["Lo siento, no tengo una respuesta para eso."])[0]

def initiate_conversation():
    topics = [
        "¿Cuál es tu libro favorito y por qué?",
        "Si pudieras viajar a cualquier lugar del mundo, ¿a dónde irías?",
        "¿Qué superpoder te gustaría tener?",
        "¿Cuál es tu película favorita y por qué?",
        "¿Qué harías si ganaras la lotería?",
        "¿Cuál es tu comida favorita?",
        "¿Qué habilidades te gustaría aprender?",
        "¿Cuál es tu deporte favorito?",
        "¿Tienes alguna mascota? Cuéntame sobre ella.",
        "¿Cuál es tu música favorita?"
    ]
    return topics.pop(0)